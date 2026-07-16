from __future__ import annotations

import io
import re

import docx as docx_lib

from . import ConversionError, ConversionResult, check_zip_safety
from .headings import promote_headings

_HEADING_STYLE = re.compile(r"^(?:Heading|제목)\s*(\d)", re.IGNORECASE)


def convert_docx(data: bytes) -> ConversionResult:
    check_zip_safety(data)
    try:
        d = docx_lib.Document(io.BytesIO(data))
    except Exception as e:
        raise ConversionError("DOCX 파일을 읽지 못했어요. 텍스트 붙여넣기로 시도해주세요.") from e

    lines: list[str] = []
    has_heading = False
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        m = _HEADING_STYLE.match(para.style.name or "")
        if m:
            has_heading = True
            lines.append("#" * min(int(m.group(1)), 6) + " " + text)
        else:
            lines.append(text)

    warnings: list[str] = []
    if d.tables:
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        warnings.append("표가 텍스트로 평탄화됐어요 — 표 안 수치 검사는 부정확할 수 있어요")

    if not lines:
        raise ConversionError("파일에서 텍스트를 찾지 못했어요. 텍스트 붙여넣기로 시도해주세요.")

    text = "\n\n".join(lines)
    if not has_heading:
        text, extra = promote_headings(text)
        warnings.extend(extra)
    return ConversionResult(text=text, warnings=warnings)
